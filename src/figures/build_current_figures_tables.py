# Part of the charls-elsa-physical-function-activity-limitations project.
# Developed for the analyses reported in the associated manuscript.
# Released under the MIT License. See LICENSE in the repository root.
# No CHARLS or ELSA individual-level data are distributed with this code.

from pathlib import Path
import os
import csv, hashlib, json, zipfile
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from matplotlib.gridspec import GridSpec
from matplotlib.ticker import NullFormatter
from matplotlib import font_manager
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(os.environ.get('FIGURE_OUTPUT_ROOT', Path(__file__).resolve().parents[2] / 'outputs' / 'figures'))
OUT.mkdir(parents=True, exist_ok=True)

available = {f.name for f in font_manager.fontManager.ttflist}
CJK = next((x for x in ['Microsoft YaHei','Microsoft YaHei UI','SimHei','Noto Sans CJK SC'] if x in available), 'DejaVu Sans')
plt.rcParams.update({'font.family':'sans-serif','font.sans-serif':[CJK,'Arial','DejaVu Sans'],
                     'axes.unicode_minus':False,'pdf.fonttype':42,'svg.fonttype':'none','font.size':9,
                     'axes.spines.top':False,'axes.spines.right':False})

# Frozen displayed values only; no statistical recomputation.
flow_charls = [
('2011年和2015年均有血液检测记录','7,648',''),
('2011年和2015年12项生物标志物均完整','7,206','排除 n=442\n至少一波存在≥1项生物标志物缺失'),
('可计算连续多系统生物标志物负担','7,200','排除 n=6\n缺少2011年性别信息，无法完成\n生物标志物负担评分'),
('生物标志物负担变化和握力变化均完整','6,155','排除 n=1,045\n缺少有效握力变化'),
('2015年无ADL/IADL依赖\n并符合主要结局风险集定义','3,036','排除 n=3,119\n2015年存在ADL/IADL依赖：n=1,124\n2015年状态无法判定：n=1,995'),
('2018年或2020年至少一波\nADL/IADL结局可判定','2,884','排除 n=152\n2018年和2020年结局均无法判定'),
('可生成至少一条有效人–区间记录','2,671','排除 n=213\n无法构建有效风险区间'),
('完全调整模型所需协变量均完整','2,499','排除 n=172\n至少1项完全调整模型协变量缺失'),
]
flow_elsa = [
('第4轮和第6轮握力变化可计算','6,143',''),
('第6轮无ADL/IADL活动困难','4,688','排除 n=1,455\n基线存在活动困难或状态无法判定'),
('第7轮或第8轮至少一波结局可判定','4,285','排除 n=403\n随访结局无法判定'),
('完全调整模型所需协变量均完整','3,991','排除 n=294\n至少1项完全调整模型协变量缺失'),
]

def draw_flow():
    fig = plt.figure(figsize=(13.2,8.2), facecolor='white')
    ax=fig.add_axes([0,0,1,1]); ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis('off')
    ax.text(.255,.962,'A  CHARLS主要分析',ha='center',va='center',fontsize=11,fontweight='bold')
    ax.text(.755,.962,'B  ELSA主要概念性重复',ha='center',va='center',fontsize=11,fontweight='bold')
    ax.plot([.5,.5],[.07,.93],color='#D0D3D6',lw=.9)
    def panel(rows,x0,x1,top,bottom):
        n=len(rows); ys=np.linspace(top,bottom,n); bw=.285; bh=.062 if n>5 else .075
        bx=x0+.045; ex=bx+bw+.018; ew=x1-ex-.025
        for i,((label,nval,excl),y) in enumerate(zip(rows,ys)):
            face='#EAF1F5' if i==n-1 else '#FAFAFA'
            box=FancyBboxPatch((bx,y-bh/2),bw,bh,boxstyle='round,pad=.004,rounding_size=.005',
                               lw=.9,edgecolor='#68737A',facecolor=face)
            ax.add_patch(box)
            ax.text(bx+bw/2,y,f'{label}\nN={nval}',ha='center',va='center',fontsize=8.2,linespacing=1.15)
            if i>0 and excl:
                ax.text(ex,y,excl,ha='left',va='center',fontsize=7.5,color='#4B5155',linespacing=1.18)
            if i<n-1:
                yn=ys[i+1]
                ax.add_patch(FancyArrowPatch((bx+bw/2,y-bh/2-.004),(bx+bw/2,yn+bh/2+.004),
                                             arrowstyle='-|>',mutation_scale=8,lw=.75,color='#777'))
    panel(flow_charls,.01,.49,.89,.18)
    panel(flow_elsa,.51,.99,.82,.34)
    ax.text(.255,.095,'主要分析：2,499人；4,449条人–区间记录；674例首次ADL/IADL依赖事件',
            ha='center',va='center',fontsize=8.7,fontweight='bold')
    ax.text(.755,.245,'主要概念性重复：3,991人；7,233条人–区间记录；733例首次ADL/IADL活动困难事件',
            ha='center',va='center',fontsize=8.7,fontweight='bold')
    base=OUT/'FIGURE_1_CHARLS_ELSA_PARTICIPANT_FLOW_ZH'
    for ext in ['png','pdf','svg']:
        fig.savefig(base.with_suffix('.'+ext),dpi=600 if ext in ('png','tiff') else None,bbox_inches='tight',pad_inches=.12,facecolor='white')
    plt.close(fig)

fig2_rows = [
('A','CHARLS首次ADL/IADL依赖：握力下降','CHARLS',1.182,1.076,1.298),
('A','CHARLS首次ADL/IADL依赖：步行时间增加','CHARLS',1.393,1.151,1.688),
('A','CHARLS首次ADL/IADL依赖：平衡恶化','CHARLS',1.408,1.130,1.754),
('A','CHARLS首次ADL/IADL依赖：五次坐站时间增加','CHARLS',1.358,1.231,1.498),
('A','ELSA首次ADL/IADL活动困难：握力下降','ELSA',1.115,1.017,1.223),
('A','ELSA首次ADL/IADL活动困难：步行时间增加','ELSA',1.516,1.335,1.722),
('A','ELSA首次ADL/IADL活动困难：平衡恶化','ELSA',1.380,1.125,1.693),
('B','CHARLS ADL依赖后完全恢复','CHARLS',0.645,0.516,0.806),
('B','CHARLS IADL依赖后完全恢复','CHARLS',0.834,0.741,0.937),
('B','ELSA ADL活动困难消失','ELSA',0.847,0.764,0.938),
('B','ELSA IADL活动困难消失','ELSA',0.865,0.771,0.969),
]

def draw_effects():
    fig=plt.figure(figsize=(12.2,8.0),facecolor='white')
    gs=GridSpec(2,3,figure=fig,width_ratios=[4.65,3.0,2.25],height_ratios=[7,4],hspace=.55,wspace=.02)
    titles={'A':'A  首次活动受限结局','B':'B  依赖后完全恢复与活动困难消失'}
    colors={'CHARLS':'#2F5D8A','ELSA':'#B56500'}
    for ix,panel in enumerate(['A','B']):
        data=[r for r in fig2_rows if r[0]==panel]; n=len(data); ys=np.arange(n)[::-1]
        axl=fig.add_subplot(gs[ix,0]); ax=fig.add_subplot(gs[ix,1],sharey=axl); axr=fig.add_subplot(gs[ix,2],sharey=axl)
        for a in [axl,axr]: a.set_xlim(0,1); a.set_ylim(-.55,n-.35); a.axis('off')
        axl.set_title(titles[panel],loc='left',fontsize=10.5,fontweight='bold',pad=12)
        axr.text(.02,n-.1,'HR（95% CI）',ha='left',va='bottom',fontsize=9,fontweight='bold')
        for y,(_,label,cohort,hr,lo,hi) in zip(ys,data):
            axl.text(.99,y,label,ha='right',va='center',fontsize=8.6)
            ax.errorbar(hr,y,xerr=[[hr-lo],[hi-hr]],fmt='s' if cohort=='CHARLS' else 'o',
                        color=colors[cohort],ecolor=colors[cohort],elinewidth=1.55,capsize=3,markersize=5.2,zorder=3)
            axr.text(.02,y,f'{hr:.3f}（{lo:.3f}–{hi:.3f}）',ha='left',va='center',fontsize=8.7)
        ax.set_xscale('log'); ax.set_xlim(.45,2.10); ax.set_ylim(-.55,n-.35); ax.set_yticks([])
        ax.axvline(1,color='#6B7280',lw=1,ls='--'); ax.grid(axis='x',color='#E5E7EB',lw=.8)
        ax.set_xticks([.5,.75,1,1.5,2]); ax.set_xticklabels(['0.50','0.75','1.00','1.50','2.00'])
        ax.xaxis.set_minor_formatter(NullFormatter()); ax.set_xlabel('风险比（HR，对数尺度）',fontsize=9)
    fig.subplots_adjust(left=.03,right=.985,top=.96,bottom=.07)
    base=OUT/'FIGURE_2_CHARLS_ELSA_DIRECTIONAL_COMPARISON_ZH'
    for ext in ['png','pdf','svg']:
        fig.savefig(base.with_suffix('.'+ext),dpi=600 if ext in ('png','tiff') else None,bbox_inches='tight',pad_inches=.12,facecolor='white')
    plt.close(fig)

table1 = [
('人口学特征',''),('年龄，岁','61.33（8.54）'),('性别',''),('　男性','1,078（43.1%）'),('　女性','1,421（56.9%）'),
('教育程度',''),('　无正规教育','1,018（40.7%）'),('　小学','682（27.3%）'),('　初中','565（22.6%）'),('　高中/职业教育','206（8.2%）'),('　大学及以上','28（1.1%）'),
('婚姻状态',''),('　已婚','2,276（91.1%）'),('　未婚/其他','223（8.9%）'),('居住地',''),('　农村/其他','2,080（83.2%）'),('　城镇','419（16.8%）'),
('健康行为与健康状况',''),('吸烟状态',''),('　从不吸烟','1,568（62.7%）'),('　既往吸烟','215（8.6%）'),('　当前吸烟','716（28.7%）'),
('饮酒状态',''),('　从不饮酒','1,533（61.3%）'),('　既往饮酒','214（8.6%）'),('　当前饮酒','752（30.1%）'),
('体重指数，kg/m²','23.84（21.47，26.51）'),('CES-D 10得分，分','8.34（6.03）'),('慢性病数量，种','1.49（1.40）'),
('多系统生物标志物负担',''),('　2011年负担','0.05（0.34）'),('　2015年负担','0.12（−0.07，0.36）'),('　2011—2015年负担变化','0.13（0.30）'),
('握力',''),('　2011年握力，kg','32.93（9.61）'),('　2015年握力，kg','30.54（9.06）'),('　2011—2015年握力下降，kg','2.39（7.40）'),
('随访与结局',''),('进入2015—2018年风险区间，n','2,499'),('2018年首次事件，n','440'),('进入2018—2020年风险区间，n','1,950'),('2020年新增首次事件，n','234'),('人–区间记录总数，n','4,449'),('累计首次事件，n','674')]

primary = [
('确认性主要分析','首次合并ADL/IADL依赖','生物标志物负担增加（每1个变化值SD）','2,499','674','1.075','0.995–1.160','0.066','—'),
('确认性主要分析','首次合并ADL/IADL依赖','握力下降（每1个变化值SD）','2,499','674','1.182','1.076–1.298','<0.001','—'),
('确认性主要分析','首次合并ADL/IADL依赖','负担增加×握力下降','2,499','674','1.042','0.960–1.131','0.329','—')]
discovery = [
('发现性扩展','ADL首次依赖','生物标志物负担增加','3,585','429','1.123','1.018–1.238','0.021','0.031'),
('发现性扩展','ADL首次依赖','握力下降','3,585','429','1.203','1.065–1.360','0.003','0.006'),
('发现性扩展','ADL首次依赖','负担增加×握力下降','3,585','429','1.030','0.950–1.118','0.469','0.563'),
('发现性扩展','IADL首次依赖','生物标志物负担增加','3,859','819','1.125','1.046–1.210','0.002','0.005'),
('发现性扩展','IADL首次依赖','握力下降','3,859','819','1.215','1.116–1.323','<0.001','<0.001'),
('发现性扩展','IADL首次依赖','负担增加×握力下降','3,859','819','1.021','0.944–1.104','0.605','0.605'),
('发现性扩展','步行领域','生物标志物负担增加','955','361','1.140','1.030–1.262','0.011','0.025'),
('发现性扩展','步行领域','步行时间增加','955','361','1.393','1.151–1.688','<0.001','0.003'),
('发现性扩展','步行领域','负担增加×步行时间增加','955','361','1.136','0.999–1.293','0.052','0.089'),
('发现性扩展','平衡领域','平衡恶化','2,522','674','1.408','1.130–1.754','0.002','0.007'),
('发现性扩展','平衡领域','生物标志物负担增加','2,522','674','1.077','0.985–1.178','0.103','0.133'),
('发现性扩展','平衡领域','负担增加×平衡恶化','2,522','674','0.931','0.752–1.152','0.511','0.511'),
('发现性扩展','五次坐站领域','生物标志物负担增加','2,443','643','1.045','0.957–1.140','0.327','0.368'),
('发现性扩展','五次坐站领域','五次坐站时间增加','2,443','643','1.358','1.231–1.498','<0.001','<0.001'),
('发现性扩展','五次坐站领域','负担增加×五次坐站时间增加','2,443','643','0.930','0.863–1.003','0.059','0.089'),
('发现性扩展','ADL完全恢复','生物标志物负担增加','242','170','1.009','0.873–1.167','0.900','0.999'),
('发现性扩展','ADL完全恢复','握力下降','242','170','0.645','0.516–0.806','<0.001','<0.001'),
('发现性扩展','ADL完全恢复','负担增加×握力下降','242','170','0.963','0.818–1.133','0.650','0.975'),
('发现性扩展','IADL完全恢复','生物标志物负担增加','823','530','1.000','0.909–1.101','0.999','0.999'),
('发现性扩展','IADL完全恢复','握力下降','823','530','0.834','0.741–0.937','0.002','0.007'),
('发现性扩展','IADL完全恢复','负担增加×握力下降','823','530','1.026','0.932–1.129','0.606','0.975')]
table2=primary+discovery

table3=[
('主要概念性重复','首次ADL/IADL活动困难','握力下降','3,991','7,233','733','1.115','1.017–1.223','0.020'),
('次要协调分析','5项ADL收到帮助','握力下降','4,804','8,999','262','1.319','1.150–1.512','<0.001'),
('次要并列生理分析','首次ADL/IADL活动困难','生物标志物负担增加','1,139','2,098','175','1.079','0.920–1.267','0.350'),
('次要并列生理分析','首次ADL/IADL活动困难','握力下降','1,139','2,098','175','1.263','1.045–1.526','0.015'),
('次要并列生理分析','首次ADL/IADL活动困难','负担增加×握力下降','1,139','2,098','175','1.016','0.876–1.178','0.833'),
('领域分析','首次ADL/IADL活动困难','步行时间增加','2,826','5,035','617','1.516','1.335–1.722','<0.001'),
('领域分析','首次ADL/IADL活动困难','平衡恶化','4,075','7,375','764','1.380','1.125–1.693','0.002'),
('状态改善','ADL活动困难消失','握力下降','776','1,164','414','0.847','0.764–0.938','0.001'),
('状态改善','IADL活动困难消失','握力下降','718','1,099','339','0.865','0.771–0.969','0.013')]

def save_csvs():
    specs=[('TABLE_1_SOURCE_DATA.csv',['特征','主要分析样本（N=2,499）'],table1),
           ('TABLE_2_SOURCE_DATA.csv',['分析层级','结局/领域','暴露项','N','事件数','HR','95% CI','p值','q值'],table2),
           ('TABLE_3_SOURCE_DATA.csv',['分析层级','结局','暴露','N','人–区间记录','事件数','HR','95% CI','p值'],table3),
           ('FIGURE_2_SOURCE_DATA.csv',['panel','label','cohort','HR','CI_low','CI_high'],fig2_rows)]
    for name,header,rows in specs:
        with (OUT/name).open('w',newline='',encoding='utf-8-sig') as f:
            w=csv.writer(f); w.writerow(header); w.writerows(rows)

def set_cell_shading(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None: shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'),fill)

def set_cell_text(cell,text,bold=False,size=8.5,align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text=''; p=cell.paragraphs[0]; p.alignment=align; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
    r=p.add_run(str(text)); r.bold=bold; r.font.size=Pt(size); r.font.name='Arial'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def repeat_header(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:tblHeader'); el.set(qn('w:val'),'true'); trPr.append(el)

def prevent_split(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); trPr.append(el)

def add_caption(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.keep_with_next=True; r=p.add_run(text); r.bold=True; r.font.size=Pt(10); r.font.name='Arial'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')

def add_note(doc,text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(8); r=p.add_run(text); r.font.size=Pt(8); r.font.name='Arial'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')

def add_table(doc,headers,rows,widths=None,group_cols=()):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; repeat_header(t.rows[0])
    for j,h in enumerate(headers): set_cell_text(t.rows[0].cells[j],h,True,8,WD_ALIGN_PARAGRAPH.CENTER); set_cell_shading(t.rows[0].cells[j],'DCE6EC')
    prev=[None]*len(headers)
    for row in rows:
        cells=t.add_row().cells; prevent_split(t.rows[-1])
        for j,val in enumerate(row):
            shown='' if j in group_cols and prev[j]==val else val
            align=WD_ALIGN_PARAGRAPH.LEFT if j<3 else WD_ALIGN_PARAGRAPH.CENTER
            set_cell_text(cells[j],shown,False,8,align)
        prev=list(row)
    if widths:
        for row in t.rows:
            for j,w in enumerate(widths): row.cells[j].width=Cm(w)
    return t

def build_docx():
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(1.0); sec.bottom_margin=Cm(1.0); sec.left_margin=Cm(1.5); sec.right_margin=Cm(1.5)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('CHARLS–ELSA主要图表（中文投稿版）'); r.bold=True; r.font.size=Pt(15); r.font.name='Arial'; r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    add_caption(doc,'图1　CHARLS与ELSA分析对象筛选流程')
    doc.add_picture(str(OUT/'FIGURE_1_CHARLS_ELSA_PARTICIPANT_FLOW_ZH.png'),width=Cm(17.5)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_note(doc,'注：Panel A为CHARLS主要分析，Panel B为ELSA主要概念性重复。CHARLS结局为需要帮助或无法完成的ADL/IADL依赖；ELSA结局为完成ADL/IADL活动存在困难，两者并不等价。')
    doc.add_page_break()
    add_caption(doc,'表1　CHARLS主要分析完整病例样本的特征（N=2,499）')
    t=add_table(doc,['特征','主要分析样本（N=2,499）'],table1,[11.5,5.2])
    for row in t.rows[1:]:
        if not row.cells[1].text:
            set_cell_shading(row.cells[0],'F1F4F6'); set_cell_shading(row.cells[1],'F1F4F6');
            for c in row.cells: c.paragraphs[0].runs[0].bold=True
    add_note(doc,'注：连续变量按原表所示以均值（标准差）或中位数（第25、第75百分位数）表示，分类变量以n（%）表示。BMI缺失14例，CES-D 10得分缺失156例；二者仅用于描述样本特征，未纳入完全调整模型。')
    doc.add_page_break()
    add_caption(doc,'表2　CHARLS确认性主要分析与发现性扩展结果')
    add_table(doc,['分析层级','结局/领域','暴露项','N','事件','HR','95% CI','p值','q值'],table2,[2.2,2.7,5.4,1.4,1.3,1.2,2.4,1.3,1.3],group_cols=(0,1))
    add_note(doc,'注：含乘积项模型中的两个主效应均为另一标准化变化暴露取0时的条件效应。确认性主要分析不进行FDR校正；发现性扩展在预先规定的家族内报告Benjamini–Hochberg校正q值。HR为风险比；CI为置信区间。')
    doc.add_page_break()
    add_caption(doc,'表3　ELSA预设概念性重复及次要协调分析结果')
    add_table(doc,['分析层级','结局','暴露','N','人–区间','事件','HR','95% CI','p值'],table3,[2.7,3.7,4.2,1.4,1.7,1.3,1.2,2.4,1.3],group_cols=(0,1))
    add_note(doc,'注：ELSA活动困难结局不等同于CHARLS依赖结局；两队列结果未进行统计合并。生物标志物模型中的主效应为另一标准化变化暴露取0时的条件效应。')
    doc.add_page_break()
    add_caption(doc,'图2　CHARLS与ELSA客观身体功能变化及状态改善结果')
    doc.add_picture(str(OUT/'FIGURE_2_CHARLS_ELSA_DIRECTIONAL_COMPARISON_ZH.png'),width=Cm(17.5)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_note(doc,'注：Panel A展示CHARLS首次ADL/IADL依赖与ELSA首次ADL/IADL活动困难结果，HR>1表示相应事件发生率较高；Panel B展示CHARLS依赖后完全恢复与ELSA活动困难消失结果，HR<1表示相应状态改善发生率较低。握力下降按每1个变化值标准差表示；步行时间增加和五次坐站时间增加按每1个相应指标的变化值标准差表示；平衡恶化为恶化者与未恶化者的比较。CHARLS依赖与ELSA活动困难并不等价，未计算跨队列汇总效应。横轴采用对数HR尺度，两个面板使用相同范围。')
    path=OUT/'CHARLS_ELSA_TABLES_AND_FIGURES_ZH_CURRENT.docx'; doc.save(path)
    with zipfile.ZipFile(path) as z: assert z.testzip() is None

def manifest():
    rows=[]
    for p in sorted(OUT.iterdir()):
        if p.is_file() and p.name!='FILE_MANIFEST_SHA256.csv':
            rows.append((p.name,p.stat().st_size,hashlib.sha256(p.read_bytes()).hexdigest()))
    with (OUT/'FILE_MANIFEST_SHA256.csv').open('w',newline='',encoding='utf-8-sig') as f:
        w=csv.writer(f); w.writerow(['file_name','size_bytes','sha256']); w.writerows(rows)

draw_flow(); draw_effects(); save_csvs(); build_docx(); manifest()
print(OUT)
