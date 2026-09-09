# Part of the charls-elsa-physical-function-activity-limitations project.
# Developed for the analyses reported in the associated manuscript.
# Released under the MIT License. See LICENSE in the repository root.
# No CHARLS or ELSA individual-level data are distributed with this code.

from pathlib import Path
import os
import runpy, csv, hashlib, zipfile
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from matplotlib.gridspec import GridSpec
from matplotlib.ticker import NullFormatter
import numpy as np
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

src = runpy.run_path(str(Path(__file__).with_name('build_current_figures_tables.py')))
OUT = Path(os.environ.get('FIGURE_OUTPUT_ROOT', Path(__file__).resolve().parents[2] / 'outputs' / 'figures'))
OUT.mkdir(parents=True, exist_ok=True)
add_table, add_caption, add_note = src['add_table'], src['add_caption'], src['add_note']
set_cell_shading = src['set_cell_shading']

flow_charls = [
('Blood-test records available in both 2011 and 2015','7,648',''),
('All 12 prespecified biomarkers complete at both waves','7,206','Excluded n=442\n≥1 biomarker missing\nat either wave'),
('Continuous multisystem biomarker burden calculable','7,200','Excluded n=6\nSex missing in 2011;\nburden score unavailable'),
('Biomarker-burden change and grip-strength change complete','6,155','Excluded n=1,045\nValid grip-strength\nchange unavailable'),
('Free of ADL/IADL dependency in 2015 and eligible\nfor the primary outcome risk set','3,036','Excluded n=3,119\nDependent: n=1,124\nStatus indeterminate: n=1,995'),
('ADL/IADL outcome ascertainable in at least one\nof 2018 or 2020','2,884','Excluded n=152\nOutcome indeterminate\nat both follow-up waves'),
('At least one valid person-period record generated','2,671','Excluded n=213\nNo valid at-risk interval'),
('Covariates complete for the fully adjusted model','2,499','Excluded n=172\n≥1 required covariate missing'),
]
flow_elsa = [
('Grip-strength change calculable at Waves 4 and 6','6,143',''),
('Free of ADL/IADL difficulty at Wave 6','4,688','Excluded n=1,455\nBaseline difficulty present\nor status indeterminate'),
('Outcome ascertainable at Wave 7 or Wave 8','4,285','Excluded n=403\nFollow-up outcome\nindeterminate'),
('Covariates complete for the fully adjusted model','3,991','Excluded n=294\n≥1 required covariate missing'),
]

def draw_flow():
    fig=plt.figure(figsize=(13.2,8.2),facecolor='white'); ax=fig.add_axes([0,0,1,1]); ax.set_xlim(0,1); ax.set_ylim(0,1); ax.axis('off')
    ax.text(.255,.962,'A  CHARLS primary analysis',ha='center',va='center',fontsize=11,fontweight='bold')
    ax.text(.755,.962,'B  ELSA primary conceptual replication',ha='center',va='center',fontsize=11,fontweight='bold')
    ax.plot([.5,.5],[.07,.93],color='#D0D3D6',lw=.9)
    def panel(rows,x0,x1,top,bottom):
        ys=np.linspace(top,bottom,len(rows)); bw=.285; bh=.062 if len(rows)>5 else .075; bx=x0+.045; ex=bx+bw+.018
        for i,((label,nval,excl),y) in enumerate(zip(rows,ys)):
            box=FancyBboxPatch((bx,y-bh/2),bw,bh,boxstyle='round,pad=.004,rounding_size=.005',lw=.9,edgecolor='#68737A',facecolor='#EAF1F5' if i==len(rows)-1 else '#FAFAFA')
            ax.add_patch(box); ax.text(bx+bw/2,y,f'{label}\nN={nval}',ha='center',va='center',fontsize=7.7,linespacing=1.12)
            if i and excl: ax.text(ex,y,excl,ha='left',va='center',fontsize=7.1,color='#4B5155',linespacing=1.15)
            if i<len(rows)-1:
                ax.add_patch(FancyArrowPatch((bx+bw/2,y-bh/2-.004),(bx+bw/2,ys[i+1]+bh/2+.004),arrowstyle='-|>',mutation_scale=8,lw=.75,color='#777'))
    panel(flow_charls,.01,.49,.89,.18); panel(flow_elsa,.51,.99,.82,.34)
    ax.text(.255,.095,'Primary analysis: 2,499 participants; 4,449 person-period records; 674 incident ADL/IADL dependency events',ha='center',fontsize=7.7,fontweight='bold')
    ax.text(.755,.245,'Primary conceptual replication: 3,991 participants; 7,233 person-period records; 733 incident ADL/IADL difficulty events',ha='center',fontsize=7.7,fontweight='bold')
    base=OUT/'FIGURE_1_CHARLS_ELSA_PARTICIPANT_FLOW_EN'
    for ext in ['png','pdf','svg']:
        fig.savefig(base.with_suffix('.'+ext),dpi=600 if ext=='png' else None,bbox_inches='tight',pad_inches=.12,facecolor='white')
    plt.close(fig)

fig2_rows=[
('A','CHARLS incident ADL/IADL dependency: grip-strength decline','CHARLS',1.182,1.076,1.298),
('A','CHARLS incident ADL/IADL dependency: increased walking time','CHARLS',1.393,1.151,1.688),
('A','CHARLS incident ADL/IADL dependency: balance deterioration','CHARLS',1.408,1.130,1.754),
('A','CHARLS incident ADL/IADL dependency: increased five-chair-rise time','CHARLS',1.358,1.231,1.498),
('A','ELSA incident ADL/IADL difficulty: grip-strength decline','ELSA',1.115,1.017,1.223),
('A','ELSA incident ADL/IADL difficulty: increased walking time','ELSA',1.516,1.335,1.722),
('A','ELSA incident ADL/IADL difficulty: balance deterioration','ELSA',1.380,1.125,1.693),
('B','CHARLS complete recovery from ADL dependency','CHARLS',.645,.516,.806),
('B','CHARLS complete recovery from IADL dependency','CHARLS',.834,.741,.937),
('B','ELSA resolution of ADL difficulty','ELSA',.847,.764,.938),
('B','ELSA resolution of IADL difficulty','ELSA',.865,.771,.969)]

def draw_effects():
    fig=plt.figure(figsize=(12.2,8),facecolor='white'); gs=GridSpec(2,3,figure=fig,width_ratios=[4.9,3,2.25],height_ratios=[7,4],hspace=.55,wspace=.02)
    titles={'A':'A  Incident activity-limitation outcomes','B':'B  Complete recovery from dependency and resolution of difficulty'}; colors={'CHARLS':'#2F5D8A','ELSA':'#B56500'}
    for ix,panel in enumerate(['A','B']):
        data=[r for r in fig2_rows if r[0]==panel]; n=len(data); ys=np.arange(n)[::-1]
        axl=fig.add_subplot(gs[ix,0]); ax=fig.add_subplot(gs[ix,1],sharey=axl); axr=fig.add_subplot(gs[ix,2],sharey=axl)
        for a in [axl,axr]: a.set_xlim(0,1); a.set_ylim(-.55,n-.35); a.axis('off')
        axl.set_title(titles[panel],loc='left',fontsize=10.2,fontweight='bold',pad=12); axr.text(.02,n-.1,'HR (95% CI)',ha='left',va='bottom',fontsize=9,fontweight='bold')
        for y,(_,label,cohort,hr,lo,hi) in zip(ys,data):
            axl.text(.99,y,label,ha='right',va='center',fontsize=8.3)
            ax.errorbar(hr,y,xerr=[[hr-lo],[hi-hr]],fmt='s' if cohort=='CHARLS' else 'o',color=colors[cohort],ecolor=colors[cohort],elinewidth=1.55,capsize=3,markersize=5.2,zorder=3)
            axr.text(.02,y,f'{hr:.3f} ({lo:.3f}–{hi:.3f})',ha='left',va='center',fontsize=8.7)
        ax.set_xscale('log'); ax.set_xlim(.45,2.10); ax.set_ylim(-.55,n-.35); ax.set_yticks([]); ax.axvline(1,color='#6B7280',lw=1,ls='--'); ax.grid(axis='x',color='#E5E7EB',lw=.8)
        ax.set_xticks([.5,.75,1,1.5,2]); ax.set_xticklabels(['0.50','0.75','1.00','1.50','2.00']); ax.xaxis.set_minor_formatter(NullFormatter()); ax.set_xlabel('Hazard ratio (log scale)',fontsize=9)
    fig.subplots_adjust(left=.025,right=.985,top=.96,bottom=.07); base=OUT/'FIGURE_2_CHARLS_ELSA_DIRECTIONAL_COMPARISON_EN'
    for ext in ['png','pdf','svg']:
        fig.savefig(base.with_suffix('.'+ext),dpi=600 if ext=='png' else None,bbox_inches='tight',pad_inches=.12,facecolor='white')
    plt.close(fig)

table1_map={'人口学特征':'Demographic characteristics','年龄，岁':'Age, years','性别':'Sex','　男性':'  Male','　女性':'  Female','教育程度':'Education','　无正规教育':'  No formal education','　小学':'  Primary school','　初中':'  Middle school','　高中/职业教育':'  High school/vocational education','　大学及以上':'  College or above','婚姻状态':'Marital status','　已婚':'  Married','　未婚/其他':'  Unmarried/other','居住地':'Residence','　农村/其他':'  Rural/other','　城镇':'  Urban','健康行为与健康状况':'Health behaviors and health status','吸烟状态':'Smoking status','　从不吸烟':'  Never smoker','　既往吸烟':'  Former smoker','　当前吸烟':'  Current smoker','饮酒状态':'Drinking status','　从不饮酒':'  Never drinker','　既往饮酒':'  Former drinker','　当前饮酒':'  Current drinker','体重指数，kg/m²':'Body mass index, kg/m²','CES-D 10得分，分':'CES-D 10 score, points','慢性病数量，种':'Number of chronic conditions','多系统生物标志物负担':'Multisystem biomarker burden','　2011年负担':'  Burden in 2011','　2015年负担':'  Burden in 2015','　2011—2015年负担变化':'  Change in burden, 2011–2015','握力':'Grip strength','　2011年握力，kg':'  Grip strength in 2011, kg','　2015年握力，kg':'  Grip strength in 2015, kg','　2011—2015年握力下降，kg':'  Grip-strength decline, 2011–2015, kg','随访与结局':'Follow-up and outcomes','进入2015—2018年风险区间，n':'Entered the 2015–2018 risk interval, n','2018年首次事件，n':'Incident events in 2018, n','进入2018—2020年风险区间，n':'Entered the 2018–2020 risk interval, n','2020年新增首次事件，n':'Additional incident events in 2020, n','人–区间记录总数，n':'Total person-period records, n','累计首次事件，n':'Total incident events, n'}
table1=[(table1_map.get(a,a),b.replace('（',' (').replace('）',')').replace('，',', ')) for a,b in src['table1']]

repls={'确认性主要分析':'Confirmatory primary analysis','发现性扩展':'Exploratory extension','首次合并ADL/IADL依赖':'Incident combined ADL/IADL dependency','ADL首次依赖':'Incident ADL dependency','IADL首次依赖':'Incident IADL dependency','步行领域':'Walking domain','平衡领域':'Balance domain','五次坐站领域':'Five-chair-rise domain','ADL完全恢复':'Complete recovery from ADL dependency','IADL完全恢复':'Complete recovery from IADL dependency','生物标志物负担增加（每1个变化值SD）':'Increase in biomarker burden (per 1 SD of change)','生物标志物负担增加':'Increase in biomarker burden','握力下降（每1个变化值SD）':'Grip-strength decline (per 1 SD of change)','握力下降':'Grip-strength decline','负担增加×握力下降':'Burden increase × grip-strength decline','步行时间增加':'Increased walking time','负担增加×步行时间增加':'Burden increase × increased walking time','平衡恶化':'Balance deterioration','负担增加×平衡恶化':'Burden increase × balance deterioration','五次坐站时间增加':'Increased five-chair-rise time','负担增加×五次坐站时间增加':'Burden increase × increased five-chair-rise time','主要概念性重复':'Primary conceptual replication','次要协调分析':'Secondary harmonized analysis','次要并列生理分析':'Secondary parallel physiological analysis','领域分析':'Domain analysis','状态改善':'State improvement','首次ADL/IADL活动困难':'Incident ADL/IADL difficulty','5项ADL收到帮助':'Receipt of assistance with five ADL items','ADL活动困难消失':'Resolution of ADL difficulty','IADL活动困难消失':'Resolution of IADL difficulty','负担增加×握力下降':'Burden increase × grip-strength decline'}
def tr(x):
    y=x
    for a in sorted(repls,key=len,reverse=True): y=y.replace(a,repls[a])
    return y
table2=[tuple(tr(str(x)) for x in row) for row in src['table2']]
table3=[tuple(tr(str(x)) for x in row) for row in src['table3']]

def build_docx():
    doc=Document(); sec=doc.sections[0]; sec.top_margin=Cm(1); sec.bottom_margin=Cm(1); sec.left_margin=Cm(1.5); sec.right_margin=Cm(1.5)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('CHARLS–ELSA Main Tables and Figures'); r.bold=True; r.font.size=Pt(15); r.font.name='Arial'
    add_caption(doc,'Figure 1. Participant flow for the CHARLS and ELSA analyses')
    doc.add_picture(str(OUT/'FIGURE_1_CHARLS_ELSA_PARTICIPANT_FLOW_EN.png'),width=Cm(17.5)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_note(doc,'Note: Panel A shows the CHARLS primary analysis and Panel B the ELSA primary conceptual replication. CHARLS measured ADL/IADL dependency (requiring assistance or being unable to perform an activity), whereas ELSA measured difficulty performing ADL/IADL activities; these outcomes are not equivalent.')
    doc.add_page_break(); add_caption(doc,'Table 1. Characteristics of the complete-case CHARLS primary-analysis sample (N=2,499)')
    t=add_table(doc,['Characteristic','Primary-analysis sample (N=2,499)'],table1,[11.5,5.2])
    for row in t.rows[1:]:
        if not row.cells[1].text:
            set_cell_shading(row.cells[0],'F1F4F6'); set_cell_shading(row.cells[1],'F1F4F6')
            for c in row.cells: c.paragraphs[0].runs[0].bold=True
    add_note(doc,'Note: Continuous variables are presented as mean (standard deviation) or median (25th, 75th percentile), as shown; categorical variables are presented as n (%). BMI was missing for 14 participants and the CES-D 10 score for 156 participants; both variables were used for descriptive purposes only and were not included in the fully adjusted model.')
    doc.add_page_break(); add_caption(doc,'Table 2. CHARLS confirmatory primary analysis and exploratory extensions')
    add_table(doc,['Analysis level','Outcome/domain','Exposure term','N','Events','HR','95% CI','p value','q value'],table2,[2.2,2.7,5.4,1.4,1.3,1.2,2.4,1.3,1.3],group_cols=(0,1))
    add_note(doc,'Note: In models containing a product term, each main-effect HR is conditional on the other standardized change exposure being 0. The confirmatory primary analysis was not subject to FDR correction; exploratory extensions report Benjamini–Hochberg-adjusted q values within the prespecified families. HR, hazard ratio; CI, confidence interval.')
    doc.add_page_break(); add_caption(doc,'Table 3. Prespecified ELSA conceptual replication and secondary harmonized analyses')
    add_table(doc,['Analysis level','Outcome','Exposure','N','Person-periods','Events','HR','95% CI','p value'],table3,[2.7,3.7,4.2,1.4,1.7,1.3,1.2,2.4,1.3],group_cols=(0,1))
    add_note(doc,'Note: The ELSA difficulty outcome is not equivalent to the CHARLS dependency outcome, and results from the two cohorts were not statistically pooled. Main effects in the biomarker model are conditional on the other standardized change exposure being 0.')
    doc.add_page_break(); add_caption(doc,'Figure 2. Objective physical-function changes and state-improvement outcomes in CHARLS and ELSA')
    doc.add_picture(str(OUT/'FIGURE_2_CHARLS_ELSA_DIRECTIONAL_COMPARISON_EN.png'),width=Cm(17.5)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
    add_note(doc,'Note: Panel A shows incident ADL/IADL dependency in CHARLS and incident ADL/IADL difficulty in ELSA; HR >1 indicates a higher event rate. Panel B shows complete recovery from dependency in CHARLS and resolution of activity difficulty in ELSA; HR <1 indicates a lower rate of state improvement. Grip-strength decline is expressed per 1 SD of change; increased walking time and five-chair-rise time are expressed per 1 SD of change in the corresponding measure; balance deterioration compares participants with versus without deterioration. CHARLS dependency and ELSA difficulty are not equivalent, and no pooled cross-cohort effect was calculated. Both panels use the same log HR scale.')
    path=OUT/'CHARLS_ELSA_TABLES_AND_FIGURES_EN_CURRENT.docx'; doc.save(path)
    with zipfile.ZipFile(path) as z: assert z.testzip() is None

def save_sources():
    for name,headers,rows in [('TABLE_1_SOURCE_DATA_EN.csv',['Characteristic','Primary-analysis sample (N=2,499)'],table1),('TABLE_2_SOURCE_DATA_EN.csv',['Analysis level','Outcome/domain','Exposure term','N','Events','HR','95% CI','p value','q value'],table2),('TABLE_3_SOURCE_DATA_EN.csv',['Analysis level','Outcome','Exposure','N','Person-periods','Events','HR','95% CI','p value'],table3)]:
        with (OUT/name).open('w',newline='',encoding='utf-8-sig') as f: w=csv.writer(f); w.writerow(headers); w.writerows(rows)

draw_flow(); draw_effects(); save_sources(); build_docx()
print(OUT)
